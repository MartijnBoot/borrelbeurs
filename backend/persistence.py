# backend/persistence.py
import os
import json
import tempfile
import datetime
from typing import Dict, List, Optional
import errno

import threading
_NEWS_LOCK = threading.Lock()

import numpy as np
from exchange.engine import ExchangeState # For type hinting
from .config import EARNINGS_DIR, STATIC_DIR, DEFAULT_CONFIG_PATH
NEWS_PATH = os.path.join(STATIC_DIR, "news.json")

# --- excel logging (with CSV fallback if openpyxl missing) ---
try:
    from openpyxl import Workbook, load_workbook
    HAVE_XLSX = True
except Exception:
    HAVE_XLSX = False
    import csv

def atomic_write_json(path: str, data):
    d = os.path.dirname(path) or "."
    os.makedirs(d, exist_ok=True)
    blob = json.dumps(data, ensure_ascii=False, indent=2)

    fd, tmp_path = tempfile.mkstemp(dir=d, prefix=".tmp_", text=True)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(blob); f.flush(); os.fsync(f.fileno())
        try:
            os.replace(tmp_path, path)  # atomic on normal files
            tmp_path = None
        except OSError as e:
            if e.errno in (errno.EBUSY, errno.EPERM):   # bind-mounted file on Windows
                with open(path, "w", encoding="utf-8") as f:
                    f.write(blob)
                os.remove(tmp_path); tmp_path = None
            else:
                raise
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except: pass

# ----------------------------- engine load/save -----------------------------
def load_engine_from_config(path: str) -> ExchangeState:
    data = {}
    try:
        if os.path.isdir(path):
            raise IsADirectoryError(f"{path} is a directory, expected file")
        if os.path.exists(path) and os.path.getsize(path) > 0:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
    except Exception as e:
        print(f"[config] could not read {path}: {e}")
    # If live config is empty/missing, fall back to the baked-in template
    if not data:
        try:
            with open(DEFAULT_CONFIG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            print(f"[config] loaded defaults from {DEFAULT_CONFIG_PATH}")
        except Exception:
            pass
    eng = ExchangeState.from_persist(data or {})
    persist = eng.to_persist()
    if "bar_price" in (data or {}):
        persist["bar_price"] = data["bar_price"]
    atomic_write_json(path, persist)
    return eng


def save_engine(path: str, eng: ExchangeState):
    atomic_write_json(path, eng.to_persist())

# ----------------------------- earnings workbook -----------------------------
current_earnings_filename: Optional[str] = None

def _ts_str() -> str:
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

def start_new_earnings_workbook() -> str:
    """Create a new earnings file for the current game and return filename."""
    global current_earnings_filename
    stamp = _ts_str()
    if HAVE_XLSX:
        current_earnings_filename = f"earnings_{stamp}.xlsx"
        path = os.path.join(EARNINGS_DIR, current_earnings_filename)
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["ts_iso","ts_ms","round","drink","qty","unit_price","total_price","p_cont","p_min","p_max"])
        wb.save(path)
    else:
        current_earnings_filename = f"earnings_{stamp}.csv"
        path = os.path.join(EARNINGS_DIR, current_earnings_filename)
        with open(path, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["ts_iso","ts_ms","round","drink","qty","unit_price","total_price","p_cont","p_min","p_max"])
    return current_earnings_filename

def append_sale_to_workbook(drink: str, qty: float, unit_price: float,
                            p_cont: float, p_min: float, p_max: float, t_round: int, ts_ms: int):
    """Append a sale to the current earnings file."""
    if not current_earnings_filename:
        start_new_earnings_workbook()
    path = os.path.join(EARNINGS_DIR, current_earnings_filename)
    ts_iso = datetime.datetime.fromtimestamp(ts_ms/1000.0).isoformat()
    total_price = unit_price * qty
    if HAVE_XLSX and current_earnings_filename.endswith(".xlsx"):
        wb = load_workbook(path)
        try:
            ws = wb["Sales"]
            ws.append([ts_iso, ts_ms, t_round, drink, qty, unit_price, total_price, p_cont, p_min, p_max])
            wb.save(path)
        finally:
            wb.close()
    else:
        with open(path, "a", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow([ts_iso, ts_ms, t_round, drink, qty, unit_price, total_price, p_cont, p_min, p_max])

def earnings_summary_from_workbook(names: List[str], p0: np.ndarray) -> Dict:
    """Compute totals + cumulative + p0 hypothetical from the current earnings file."""
    global current_earnings_filename
    if not current_earnings_filename:
        return {"total": 0.0, "per_drink": {nm: 0.0 for nm in names}, "count": 0,
                "series": [], "file_url": None, "total_qty": 0,
                "p0_total": 0.0, "p0_per_drink": {nm: 0.0 for nm in names}}

    path = os.path.join(EARNINGS_DIR, current_earnings_filename)
    if not os.path.exists(path):
        return {"total": 0.0, "per_drink": {nm: 0.0 for nm in names}, "count": 0,
                "series": [], "file_url": None, "total_qty": 0,
                "p0_total": 0.0, "p0_per_drink": {nm: 0.0 for nm in names}}

    per = {nm: 0.0 for nm in names}
    per_p0 = {nm: 0.0 for nm in names}
    total = 0.0
    total_p0 = 0.0
    series, count, cum, total_qty = [], 0, 0.0, 0.0
    p0_map = {nm: float(p0[i]) for i, nm in enumerate(names)}

    if HAVE_XLSX and current_earnings_filename.endswith(".xlsx"):
        wb = load_workbook(path, read_only=True)
        try:
            ws = wb["Sales"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                ts_ms, drink, qty, unit_price = int(row[1]), str(row[3]), float(row[4]), float(row[5])
                val = qty * unit_price
                total += val
                per[drink] = per.get(drink, 0.0) + val
                p0_val = qty * p0_map.get(drink, 0.0)
                total_p0 += p0_val
                per_p0[drink] = per_p0.get(drink, 0.0) + p0_val
                total_qty += qty
                cum += val
                series.append({"ts": ts_ms, "cum": round(cum, 2)})
                count += 1
        finally:
            wb.close()
    else:
        with open(path, newline="", encoding="utf-8") as f:
            r = csv.DictReader(f)
            for row in r:
                ts_ms = int(row["ts_ms"]); drink = row["drink"]
                qty = float(row["qty"]); unit_price = float(row["unit_price"])
                val = qty * unit_price
                total += val
                per[drink] = per.get(drink, 0.0) + val
                p0_val = qty * p0_map.get(drink, 0.0)
                total_p0 += p0_val
                per_p0[drink] = per_p0.get(drink, 0.0) + p0_val
                total_qty += qty
                cum += val
                series.append({"ts": ts_ms, "cum": round(cum, 2)})
                count += 1

    file_url = f"/static/earnings/{current_earnings_filename}"
    return {"total": round(total, 2),
            "per_drink": {k: round(v, 2) for k, v in per.items()},
            "count": count,
            "series": series,
            "file_url": file_url,
            "total_qty": int(total_qty),
            "p0_total": round(total_p0, 2),
            "p0_per_drink": {k: round(v, 2) for k, v in per_p0.items()}}

def finalize_current_earnings(names: List[str], p0: np.ndarray) -> Optional[str]:
    """Write a 'Summary' with totals into the current earnings file."""
    global current_earnings_filename
    if not current_earnings_filename:
        return None
    path = os.path.join(EARNINGS_DIR, current_earnings_filename)
    if not os.path.exists(path):
        return None
    
    summary = earnings_summary_from_workbook(names, p0)
    stamp = datetime.datetime.now().isoformat(timespec="seconds")

    if HAVE_XLSX and current_earnings_filename.endswith(".xlsx"):
        wb = load_workbook(path)
        try:
            if "Summary" in wb.sheetnames:
                wb.remove(wb["Summary"])
            ws = wb.create_sheet("Summary")
            ws.append(["finalized_at", stamp])
            ws.append(["total_revenue", summary["total"]])
            ws.append(["total_qty", summary["total_qty"]])
            ws.append(["p0_total", summary["p0_total"]])
            ws.append([])
            ws.append(["Per drink (actual)"])
            ws.append(["drink", "revenue"])
            for nm in names:
                ws.append([nm, summary["per_drink"].get(nm, 0.0)])
            ws.append([])
            ws.append(["Per drink (p0 hypothetical)"])
            ws.append(["drink", "revenue_p0"])
            for nm in names:
                ws.append([nm, summary["p0_per_drink"].get(nm, 0.0)])
            wb.save(path)
        finally:
            wb.close()
        return path
    else:
        base = os.path.splitext(path)[0]
        spath = base + "_summary.csv"
        with open(spath, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["finalized_at", stamp])
            w.writerow(["total_revenue", summary["total"]])
            w.writerow(["total_qty", summary["total_qty"]])
            w.writerow(["p0_total", summary["p0_total"]])
            w.writerow([])
            w.writerow(["drink", "revenue", "revenue_p0"])
            for nm in names:
                w.writerow([nm, summary["per_drink"].get(nm, 0.0), summary["p0_per_drink"].get(nm, 0.0)])
        return spath

# ----------------------------- news store -----------------------------
def _read_news_list():
    if not os.path.exists(NEWS_PATH):
        return []
    try:
        with open(NEWS_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return []

def get_news_list():
    with _NEWS_LOCK:
        data = _read_news_list()
        data.sort(key=lambda x: x.get("ts_ms", 0))
        return data

def append_news_item(item: Dict):
    with _NEWS_LOCK:
        data = _read_news_list()
        data.append(item)
        atomic_write_json(NEWS_PATH, data)

def delete_news_item(item_id: str) -> bool:
    with _NEWS_LOCK:
        data = _read_news_list()
        new = [x for x in data if str(x.get("id")) != str(item_id)]
        if len(new) == len(data):
            return False
        atomic_write_json(NEWS_PATH, new)
        return True
