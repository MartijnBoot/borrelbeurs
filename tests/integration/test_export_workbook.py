"""The export's snapshot and workbook (Phase 7 T10: AC13, AC18, AC19, AC22, AC23; SD8, SD13, SD15).

Runs are seeded straight into the tables at chosen times and prices: the
snapshot reads nothing else, so no engine state is needed. Each workbook is read
back with `openpyxl`, and every euro cell is compared as `Decimal(str(v)) * 100`
against integer cents from SQL (R6).
"""

from __future__ import annotations

import asyncio
import io
from collections.abc import Awaitable, Callable, Sequence
from decimal import Decimal
from typing import Any, TypeVar

import openpyxl
import pytest
from sqlalchemy import insert, text, update
from sqlalchemy.ext.asyncio import AsyncConnection, AsyncEngine

from app.core.config import Settings
from app.db import models, wallets
from app.db.exports import ExportSnapshot, read_snapshot
from app.db.runs import RunNotFound, add_drink, create_draft_run
from app.db.session import create_engine, create_heavy_read_engine
from app.jobs.export_xlsx import PER_DRINK_HEADERS, SALES_HEADERS, build_workbook
from exchange import Params

T0 = 1_759_312_800_000  # 2025-10-01T10:00:00Z
GENERATED = T0 + 3_600_000

T = TypeVar("T")


def _run_async(
    settings: Settings, url: str, body: Callable[[AsyncEngine, AsyncEngine], Awaitable[T]]
) -> T:
    """`body(engine, heavy)` on both engines over the scratch database."""

    async def scenario() -> T:
        scoped = settings.model_copy(update={"database_url": url})
        engine, heavy = create_engine(scoped), create_heavy_read_engine(scoped)
        try:
            return await body(engine, heavy)
        finally:
            await engine.dispose()
            await heavy.dispose()

    return asyncio.run(scenario())


async def _drink(
    conn: AsyncConnection, run_id: int, name: str, slot: int, *, bar_price_cents: int = 250
) -> int:
    return int(
        await add_drink(
            conn,
            run_id,
            name=name,
            slot=slot,
            p_min_cents=120,
            p0_cents=250,
            p_max_cents=480,
            a=1.0,
            d=0.1,
            s0=1.0,
            c=0.1,
            bar_price_cents=bar_price_cents,
        )
    )


async def _run(conn: AsyncConnection, name: str = "Borrel") -> int:
    run_id = await create_draft_run(conn, name=name, params=Params(step_quant=0.1), run_seed=5)
    await conn.execute(
        insert(models.PriceTick),
        {"run_id": run_id, "version": 0, "source": "tick", "prices": {}, "wall_ts_ms": T0},
    )
    return run_id


Line = tuple[int, int, int, int | None]  # (drink_id, qty, unit_price_cents, bar_price_cents)


async def _order(
    conn: AsyncConnection, run_id: int, version: int, ts_ms: int, lines: Sequence[Line]
) -> int:
    order_id = (
        await conn.execute(
            insert(models.Order).returning(models.Order.order_id),
            {
                "run_id": run_id,
                "idempotency_key": f"k-export-{run_id}-{version}",
                "version": version,
                "wall_ts_ms": ts_ms,
            },
        )
    ).scalar_one()
    await conn.execute(
        insert(models.OrderLine),
        [
            {
                "order_id": order_id,
                "drink_id": drink_id,
                "qty": qty,
                "unit_price_cents": price,
                "line_total_cents": qty * price,
                "p_cont": price / 100,
                "bar_price_cents": bar,
            }
            for drink_id, qty, price, bar in lines
        ],
    )
    return int(order_id)


async def _sql_totals(engine: AsyncEngine, run_id: int) -> tuple[int, int, int]:
    """`(sum(line_total_cents), sum of qty * bar price with SD13's fallback, line count)`."""
    async with engine.connect() as conn:
        row = (
            await conn.execute(
                text(
                    "SELECT coalesce(sum(l.line_total_cents), 0),"
                    " coalesce(sum(l.qty * coalesce(l.bar_price_cents, d.bar_price_cents)), 0),"
                    " count(*)"
                    ' FROM order_line l JOIN "order" o USING (order_id)'
                    " JOIN drink d ON d.drink_id = l.drink_id WHERE o.run_id = :r"
                ),
                {"r": run_id},
            )
        ).one()
    return int(row[0]), int(row[1]), int(row[2])


async def _snapshot(heavy: AsyncEngine, run_id: int) -> ExportSnapshot:
    async with heavy.connect() as conn:
        return await read_snapshot(conn, run_id)


def _book(data: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(data))


def _cents(value: Any) -> int:
    cents = Decimal(str(value)) * 100
    assert cents == cents.to_integral_value(), value
    return int(cents)


def _rows(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows(values_only=True)]


def _summary(book: Any) -> dict[str, Any]:
    return {label: value for label, value in _rows(book["Summary"])}


def test_the_totals_equal_the_database_to_the_cent(settings: Settings, database_url: str) -> None:
    """AC18: Summary, Σ Sales and the snapshot all equal SQL; the bar price is per sale."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[Any, ...]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0, bar_price_cents=250)
            wijn = await _drink(conn, run_id, "Wijn", 1, bar_price_cents=300)
            await _order(conn, run_id, 1, T0 + 1_000, [(bier, 2, 233, 250), (wijn, 1, 317, 300)])
            await _order(conn, run_id, 2, T0 + 61_000, [(bier, 3, 241, 250)])
            # The bar price changed mid-run: the later line carries the new one (AC12).
            await _order(conn, run_id, 3, T0 + 125_000, [(wijn, 4, 299, 350)])
        snapshot = await _snapshot(heavy, run_id)
        return snapshot, await _sql_totals(engine, run_id)

    snapshot, (total, bar_total, count) = _run_async(settings, database_url, body)
    book = _book(build_workbook(snapshot, generated_at_ms=GENERATED))
    sales = _rows(book["Sales"])
    summary = _summary(book)

    assert sales[0] == SALES_HEADERS
    assert snapshot.total_cents == total
    assert snapshot.line_count == count == len(sales) - 1
    assert sum(_cents(row[5]) for row in sales[1:]) == total
    assert _cents(summary["total_revenue"]) == total
    assert _cents(summary["bar_price_total"]) == bar_total == 2 * 250 + 300 + 3 * 250 + 4 * 350
    assert summary["total_qty"] == 10
    assert summary["line_count"] == count
    assert summary["run"] == "Borrel"
    assert summary["started"] == "2025-10-01T10:00:00.000+00:00"
    assert summary["generated_at"] == "2025-10-01T11:00:00.000+00:00"
    per_drink = {row[0]: row[1:] for row in _rows(book["Per drink"])[1:]}
    assert {name: (qty, _cents(revenue)) for name, (qty, revenue, *_) in per_drink.items()} == {
        "Bier": (5, 2 * 233 + 3 * 241),
        "Wijn": (5, 317 + 4 * 299),
    }
    assert _cents(per_drink["Wijn"][2]) == 300 + 4 * 350
    assert _cents(per_drink["Bier"][3]) == round((2 * 233 + 3 * 241) / 5)


def test_a_sales_row_is_sd8s_columns_in_utc(settings: Settings, database_url: str) -> None:
    """SD8's column order, SD15's UTC stamp, the drink's name and bounds at export."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[ExportSnapshot, int, int]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0)
            order_id = await _order(conn, run_id, 7, T0 + 1_234, [(bier, 2, 233, 260)])
            await conn.execute(
                update(models.Drink)
                .where(models.Drink.drink_id == bier)
                .values(name="Pils", p_max_cents=510)
            )
        return await _snapshot(heavy, run_id), bier, order_id

    snapshot, bier, order_id = _run_async(settings, database_url, body)
    row = _rows(_book(build_workbook(snapshot, generated_at_ms=GENERATED))["Sales"])[1]

    assert row[:4] == ("2025-10-01T10:00:01.234+00:00", T0 + 1_234, "Pils", 2)
    assert [_cents(v) for v in (row[4], row[5])] == [233, 466]
    assert row[6] == pytest.approx(2.33)
    assert [_cents(v) for v in (row[7], row[8])] == [120, 510]
    assert row[9:11] == (order_id, bier)
    assert _cents(row[11]) == 260
    assert row[12] == 7


def test_added_and_removed_drinks_are_one_workbook(settings: Settings, database_url: str) -> None:
    """AC19: the removed drink's lines stay; the drink added later is in the same file."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[ExportSnapshot, int]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0)
            wijn = await _drink(conn, run_id, "Wijn", 1)
            await _order(conn, run_id, 1, T0 + 1_000, [(bier, 1, 250, 250), (wijn, 2, 260, 250)])
            await conn.execute(
                text("UPDATE drink SET removed_at = now() WHERE drink_id = :d"), {"d": wijn}
            )
            fris = await _drink(conn, run_id, "Fris", 2)
            await _order(conn, run_id, 2, T0 + 90_000, [(fris, 3, 150, 250)])
        return await _snapshot(heavy, run_id), wijn

    snapshot, wijn = _run_async(settings, database_url, body)
    book = _book(build_workbook(snapshot, generated_at_ms=GENERATED))
    sales = _rows(book["Sales"])[1:]

    assert [row[2] for row in sales] == ["Bier", "Wijn", "Fris"]
    assert [row[10] for row in sales if row[2] == "Wijn"] == [wijn]
    assert [row[0] for row in _rows(book["Per drink"])[1:]] == ["Bier", "Wijn", "Fris"]


def test_a_run_without_orders_has_headers_and_zero_totals(
    settings: Settings, database_url: str
) -> None:
    """AC22."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> ExportSnapshot:
        async with engine.begin() as conn:
            run_id = await create_draft_run(
                conn, name="Leeg", params=Params(step_quant=0.1), run_seed=1
            )
            await _drink(conn, run_id, "Bier", 0)
        return await _snapshot(heavy, run_id)

    snapshot = _run_async(settings, database_url, body)
    book = _book(build_workbook(snapshot, generated_at_ms=GENERATED))
    summary = _summary(book)

    assert book.sheetnames == ["Sales", "Per drink", "Summary"]
    assert _rows(book["Sales"]) == [SALES_HEADERS]
    assert _rows(book["Per drink"]) == [PER_DRINK_HEADERS, ("Bier", 0, 0, 0, None)]
    assert (snapshot.line_count, snapshot.total_cents) == (0, 0)
    assert summary["total_revenue"] == summary["bar_price_total"] == summary["total_qty"] == 0
    assert summary["line_count"] == 0
    assert summary["started"] is None and summary["ended"] is None


def test_an_order_committed_after_the_snapshot_opened_is_absent(
    settings: Settings, database_url: str
) -> None:
    """AC23: one REPEATABLE READ snapshot covers the lines, the count and the total."""

    class CommitAfterFirstRead:
        """The heavy connection, but a second connection commits an order after its first read."""

        def __init__(self, conn: AsyncConnection, commit: Callable[[], Awaitable[None]]) -> None:
            self._conn, self._commit, self._fired = conn, commit, False

        async def execute(self, *args: Any, **kwargs: Any) -> Any:
            result = await self._conn.execute(*args, **kwargs)
            if not self._fired:
                self._fired = True
                await self._commit()
            return result

        def __getattr__(self, name: str) -> Any:
            return getattr(self._conn, name)

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[ExportSnapshot, int, int]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0)
            await _order(conn, run_id, 1, T0 + 1_000, [(bier, 1, 250, 250)])

        async def late_order() -> None:
            async with engine.begin() as other:
                await _order(other, run_id, 2, T0 + 2_000, [(bier, 4, 260, 250)])

        async with heavy.connect() as conn:
            snapshot = await read_snapshot(CommitAfterFirstRead(conn, late_order), run_id)  # type: ignore[arg-type]
        total, _, count = await _sql_totals(engine, run_id)
        return snapshot, total, count

    snapshot, total, count = _run_async(settings, database_url, body)

    assert count == 2 and total == 250 + 4 * 260  # the late order did commit
    assert [line.version for line in snapshot.lines] == [1]
    assert (snapshot.line_count, snapshot.total_cents) == (1, 250)


def test_rows_from_a_rolled_back_image_fall_back(settings: Settings, database_url: str) -> None:
    """SD13, R14: a NULL bar price takes the drink's; a NULL product groups by `name_key`."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[ExportSnapshot, int]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0, bar_price_cents=275)
            await _order(conn, run_id, 1, T0 + 1_000, [(bier, 2, 250, None)])
            await conn.execute(
                update(models.Drink).where(models.Drink.drink_id == bier).values(product_id=None)
            )
        _, bar_total, _ = await _sql_totals(engine, run_id)
        return await _snapshot(heavy, run_id), bar_total

    snapshot, bar_total = _run_async(settings, database_url, body)
    book = _book(build_workbook(snapshot, generated_at_ms=GENERATED))

    assert snapshot.drinks[0].product_key == "name:bier"
    assert _cents(_rows(book["Sales"])[1][11]) == 275
    assert _cents(_summary(book)["bar_price_total"]) == bar_total == 550
    assert _rows(book["Per drink"])[1][:2] == ("Bier", 2)


def test_an_unknown_run_is_run_not_found(settings: Settings, database_url: str) -> None:
    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> None:
        with pytest.raises(RunNotFound):
            await _snapshot(heavy, 999_999)

    _run_async(settings, database_url, body)


def test_no_cell_names_p0_and_text_is_never_a_formula(
    settings: Settings, database_url: str
) -> None:
    """AC13 (workbook half); a drink named like a formula is written as text."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> ExportSnapshot:
        async with engine.begin() as conn:
            run_id = await _run(conn, name="=HYPERLINK(1)")
            drink = await _drink(conn, run_id, "=1+1", 0)
            await _order(conn, run_id, 1, T0 + 1_000, [(drink, 1, 250, 250)])
        return await _snapshot(heavy, run_id)

    snapshot = _run_async(settings, database_url, body)
    data = build_workbook(snapshot, generated_at_ms=GENERATED)
    book = _book(data)

    cells = [cell for sheet in book.worksheets for row in sheet.iter_rows() for cell in row]
    assert not [c.value for c in cells if isinstance(c.value, str) and "p0_" in c.value]
    assert {c.data_type for c in cells if c.value in {"=1+1", "=HYPERLINK(1)"}} == {"s"}
    assert build_workbook(snapshot, generated_at_ms=GENERATED) == data  # pure: same in, same out


def test_wallets_leave_the_workbook_unchanged(settings: Settings, database_url: str) -> None:
    """B-05 AC35: no sheet, column or value from a wallet, its entries or its settings.

    B-06 T18 (AC36): nor from a payment naming an order, or its refund."""

    async def body(engine: AsyncEngine, heavy: AsyncEngine) -> tuple[ExportSnapshot, ...]:
        async with engine.begin() as conn:
            run_id = await _run(conn)
            bier = await _drink(conn, run_id, "Bier", 0)
            paid = await _order(conn, run_id, 1, T0 + 1_000, [(bier, 2, 233, 250)])
            kept = await _order(conn, run_id, 2, T0 + 2_000, [(bier, 1, 240, 250)])
        without = await _snapshot(heavy, run_id)
        async with engine.begin() as conn:
            jan = await wallets.create_wallet(conn, name="Jan", cookie_hash="h-jan", now_ms=T0)
            topup = await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="topup",
                amount_cents=1500,
                now_ms=T0,
                run_id=run_id,
            )
            await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="reversal",
                amount_cents=-1500,
                now_ms=T0,
                run_id=run_id,
                reverses_entry_id=topup.entry_id,
            )
            await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="topup",
                amount_cents=2000,
                now_ms=T0,
                run_id=run_id,
            )
            payment = await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="payment",
                amount_cents=-466,
                now_ms=T0 + 1_000,
                run_id=run_id,
                order_id=paid,
            )
            await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="refund",
                amount_cents=466,
                now_ms=T0 + 1_500,
                run_id=run_id,
                refunds_entry_id=payment.entry_id,
            )
            await wallets.insert_entry(
                conn,
                wallet_id=jan.wallet_id,
                kind="payment",
                amount_cents=-240,
                now_ms=T0 + 2_000,
                run_id=run_id,
                order_id=kept,
            )
            await wallets.put_settings(conn, wallets.WalletSettingsRow("Borrel", 7, 3))
        return without, await _snapshot(heavy, run_id)

    without, with_wallets = _run_async(settings, database_url, body)
    before = _book(build_workbook(without, generated_at_ms=GENERATED))
    after = _book(build_workbook(with_wallets, generated_at_ms=GENERATED))

    assert with_wallets == without
    assert after.sheetnames == before.sheetnames
    for name in before.sheetnames:
        assert _rows(after[name]) == _rows(before[name]), name
